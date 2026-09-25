"""Official BME Continuous Market daily equity prices for BIAP Global.

The public BME Daily Bulletin separates equities, ETFs and certificates into
independent worksheets. BIAP reads only EQ Valores and joins it to the
authoritative FIRDS Madrid universe by ISIN, so a vendor ticker mismatch cannot
silently introduce an instrument outside the selected regulated-equity market.
"""
from __future__ import annotations

from datetime import date, datetime
import io
from typing import Iterable, Optional

import openpyxl
import requests

from .country_packs import ExchangeSpec
from .models import GlobalCompany
from .providers import GlobalProviderError

_URL = "https://www.bolsasymercados.es/dam/descargas/exchange/renta-variable/bdiario.xlsx"
_USER_AGENT = "Mozilla/5.0 BIAP-Global official-BME"
_PROVIDER_ID = "official-bme-continuous-market-daily-bulletin"


def _number(value) -> Optional[float]:
    if value in (None, ""):
        return None
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    return None if number != number else number


def _valid_isin(value) -> Optional[str]:
    text = str(value or "").strip().upper()
    return text if len(text) == 12 and text.isalnum() else None


def _date_text(value) -> Optional[str]:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = str(value or "").strip()
    if not text:
        return None
    try:
        return date.fromisoformat(text[:10]).isoformat()
    except ValueError:
        return None


class BMEOfficialDailyClient:
    provider_id = _PROVIDER_ID

    def __init__(self, *, timeout: float = 45.0) -> None:
        self.timeout = max(10.0, float(timeout))

    @staticmethod
    def supported(country: str, exchange: str) -> bool:
        return (country.upper(), exchange.upper()) == ("ES", "BME_MADRID")

    def _download(self) -> bytes:
        try:
            response = requests.get(
                _URL,
                headers={"User-Agent": _USER_AGENT, "Accept": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet,*/*"},
                timeout=self.timeout,
            )
            response.raise_for_status()
        except requests.RequestException as exc:
            raise GlobalProviderError(f"BME daily bulletin download failed: {type(exc).__name__}") from exc
        if len(response.content) < 10_000:
            raise GlobalProviderError("BME daily bulletin returned an undersized XLSX payload")
        return response.content

    @staticmethod
    def parse_workbook(content: bytes) -> list[dict]:
        try:
            workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except Exception as exc:
            raise GlobalProviderError(f"BME daily bulletin XLSX parse failed: {type(exc).__name__}") from exc
        try:
            if "EQ Valores" not in workbook.sheetnames:
                raise GlobalProviderError("BME daily bulletin has no EQ Valores worksheet")
            sheet = workbook["EQ Valores"]
            header: Optional[dict[str, int]] = None
            rows: list[dict] = []
            for values in sheet.iter_rows(values_only=True):
                cells = list(values)
                normalized = [str(value or "").strip() for value in cells]
                if header is None:
                    if "Date" in normalized and "Ticker" in normalized and "ISIN Code" in normalized and "Closing Price" in normalized:
                        header = {name: normalized.index(name) for name in (
                            "Date", "Ticker", "ISIN Code", "Security Name", "Shares Outstanding",
                            "Opening Price", "High Price", "Low Price", "Closing Price",
                            "Trades", "Shares", "Turnover",
                        ) if name in normalized}
                    continue

                def cell(name: str):
                    idx = header.get(name) if header else None
                    return cells[idx] if idx is not None and idx < len(cells) else None

                isin = _valid_isin(cell("ISIN Code"))
                ticker = str(cell("Ticker") or "").strip().upper()
                close = _number(cell("Closing Price"))
                quote_date = _date_text(cell("Date"))
                if not isin or not ticker or close is None or close <= 0 or not quote_date:
                    continue
                rows.append({
                    "ticker": ticker,
                    "isin": isin,
                    "name": str(cell("Security Name") or ticker).strip(),
                    "quoteDate": quote_date,
                    "close": close,
                    "open": _number(cell("Opening Price")),
                    "high": _number(cell("High Price")),
                    "low": _number(cell("Low Price")),
                    "trades": int(_number(cell("Trades")) or 0),
                    "shares": max(0.0, float(_number(cell("Shares")) or 0.0)),
                    "turnover": max(0.0, float(_number(cell("Turnover")) or 0.0)),
                    "sharesOutstanding": _number(cell("Shares Outstanding")),
                })
            if header is None:
                raise GlobalProviderError("BME daily bulletin equity header was not found")
            if not rows:
                raise GlobalProviderError("BME daily bulletin normalized no equity rows")
            return rows
        finally:
            workbook.close()

    def rows(self) -> list[dict]:
        return self.parse_workbook(self._download())

    def batch_quotes(
        self,
        instruments: Iterable[GlobalCompany],
        country: str,
        spec: ExchangeSpec,
    ) -> tuple[list[dict], list[str], str]:
        if not self.supported(country, spec.code):
            raise GlobalProviderError(f"BME official daily market data is not configured for {country}/{spec.code}")
        rows = self.rows()
        by_isin: dict[str, list[dict]] = {}
        for row in rows:
            by_isin.setdefault(str(row["isin"]), []).append(row)

        quotes: list[dict] = []
        for company in instruments:
            isin = _valid_isin(company.isin)
            matches = by_isin.get(isin, []) if isin else []
            if len(matches) != 1:
                continue
            row = matches[0]
            price = float(row["close"])
            volume = max(0.0, float(row.get("shares") or 0.0))
            high = _number(row.get("high"))
            low = _number(row.get("low"))
            range_position = None
            if high is not None and low is not None and high > low:
                range_position = max(0.0, min(1.0, (price - low) / (high - low)))
            quotes.append({
                "ticker": company.ticker.upper(),
                "price": price,
                "averageVolume": volume,
                "liquidityValue": max(0.0, float(row.get("turnover") or (price * volume))),
                "rangePosition": range_position,
                "quoteDate": row.get("quoteDate"),
                "mic": company.mic_code or spec.mic,
                "provider": self.provider_id,
                "sourceUrl": _URL,
                "isin": isin,
                "directoryTicker": row.get("ticker"),
                "trades": row.get("trades"),
            })

        errors: list[str] = []
        if not quotes:
            errors.append("BME official continuous-market bulletin returned no usable FIRDS-universe matches")
        latest = max((str(row.get("quoteDate") or "") for row in rows), default="")
        return quotes, errors, f"BME official Continuous Market Daily Bulletin ({latest or 'latest'})"
