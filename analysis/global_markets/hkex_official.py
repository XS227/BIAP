"""Official Hong Kong Exchange ordinary-equity universe.

HKEX publishes the complete ListOfSecurities.xlsx workbook. BIAP keeps only
primary HKD equity counters on the Main Board and GEM. Debt, ETPs, warrants,
CBBCs, REITs, investment companies, depositary receipts, trading-only lines and
RMB duplicate counters are excluded from the stock-ranking universe.
"""
from __future__ import annotations

from datetime import datetime, timezone
import io
import re
from typing import Iterable, Optional

import openpyxl
import requests

from .models import GlobalCompany, SourceEvidence
from .providers import GlobalProviderError, InstrumentUniverseProvider


_SOURCE_URL = "https://www.hkex.com.hk/eng/services/trading/securities/securitieslists/ListOfSecurities.xlsx"
_PAGE_URL = "https://www.hkex.com.hk/Services/Trading/Securities/Securities-Lists?sc_lang=en"
_PROVIDER_ID = "official-hkex-list-of-securities"
_USER_AGENT = "Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 Chrome/140 Safari/537.36 BIAP-Global"
_ALLOWED_SUBCATEGORIES = {
    "Equity Securities (Main Board)",
    "Equity Securities (GEM)",
}
_PREFERENCE_NAME = re.compile(r"(?:\bPREF\b|PREFERENCE|PREFERRED)", re.IGNORECASE)


def _text(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "").replace("\n", " ")).strip()


def _lot(value: object) -> Optional[int]:
    text = _text(value).replace(",", "")
    if not text:
        return None
    try:
        number = int(float(text))
    except (TypeError, ValueError):
        return None
    return number if number > 0 else None


def _ticker(stock_code: object) -> str:
    raw = re.sub(r"\D", "", _text(stock_code))
    if not raw:
        return ""
    value = int(raw)
    return str(value).zfill(4) if value <= 9999 else str(value).zfill(5)


def _isin(value: object) -> Optional[str]:
    text = _text(value).upper()
    return text if len(text) == 12 and text.isalnum() else None


def parse_hkex_securities_rows(rows: Iterable[Iterable[object]]) -> list[GlobalCompany]:
    materialized = [list(row) for row in rows]
    header_index: Optional[int] = None
    headers: dict[str, int] = {}
    for index, row in enumerate(materialized):
        candidate = {_text(value): pos for pos, value in enumerate(row) if _text(value)}
        if "Stock Code" in candidate and "Name of Securities" in candidate and "Category" in candidate:
            header_index = index
            headers = candidate
            break
    if header_index is None:
        raise GlobalProviderError("HKEX securities workbook header was not found")

    def get(row: list[object], name: str) -> str:
        index = headers.get(name)
        return _text(row[index]) if index is not None and index < len(row) else ""

    observed = datetime.now(timezone.utc).isoformat()
    result: list[GlobalCompany] = []
    seen: set[str] = set()
    for row in materialized[header_index + 1:]:
        category = get(row, "Category")
        subcategory = get(row, "Sub-Category")
        currency = get(row, "Trading Currency").upper()
        raw_code = get(row, "Stock Code")
        name = get(row, "Name of Securities")
        if category != "Equity" or subcategory not in _ALLOWED_SUBCATEGORIES:
            continue
        # HKEX's dual-counter model gives selected issuers a separate RMB code.
        # Ranking keeps the primary HKD counter only so one issuer is not scored
        # twice and stage-one liquidity remains in one comparable currency.
        if currency != "HKD":
            continue
        if not name or _PREFERENCE_NAME.search(name):
            continue
        ticker = _ticker(raw_code)
        if not ticker or ticker in seen:
            continue
        seen.add(ticker)
        isin = _isin(get(row, "ISIN"))
        result.append(GlobalCompany(
            country="HK",
            exchange="HKEX",
            currency="HKD",
            ticker=ticker,
            name=name,
            mic_code="XHKG",
            isin=isin,
            instrument_type="Common Stock",
            lot_size=_lot(get(row, "Board Lot")),
            raw_provider_fields={
                "official_universe": True,
                "trusted_official_equity": True,
                "hkex_stock_code": raw_code,
                "hkex_category": category,
                "hkex_subcategory": subcategory,
                "hkex_rmb_counter": get(row, "RMB Counter") or None,
                "hkex_stamp_duty": get(row, "Subject to Stamp Duty") or None,
                "hkex_shortsell_eligible": get(row, "Shortsell Eligible") or None,
                "primary_counter_scope": "HKD primary equity counter",
            },
            sources=[SourceEvidence(
                provider=_PROVIDER_ID,
                source_type="official_exchange_universe",
                source_id=f"HKEX:{raw_code or ticker}",
                source_url=_SOURCE_URL,
                observed_at=observed,
                quality=1.0,
                notes=(
                    "HKEX official List of Securities; Main Board/GEM equity, "
                    "primary HKD counter only; non-equity products and duplicate RMB counters excluded."
                ),
            )],
        ))
    if not result:
        raise GlobalProviderError("HKEX workbook returned no eligible primary ordinary equities")
    return result


class HKEXOfficialUniverseProvider(InstrumentUniverseProvider):
    provider_id = _PROVIDER_ID
    source_url = _SOURCE_URL

    def __init__(self, *, timeout: float = 45.0) -> None:
        self.timeout = max(10.0, float(timeout))
        self.last_metadata: dict = {}

    @staticmethod
    def supported(country: str, exchange: str) -> bool:
        return (country.upper(), exchange.upper()) == ("HK", "HKEX")

    def _download(self) -> bytes:
        try:
            response = requests.get(
                self.source_url,
                headers={
                    "User-Agent": _USER_AGENT,
                    "Accept": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet,*/*",
                    "Referer": _PAGE_URL,
                    "Accept-Language": "en-US,en;q=0.9",
                },
                timeout=self.timeout,
            )
            response.raise_for_status()
        except requests.RequestException as exc:
            raise GlobalProviderError(f"HKEX securities workbook request failed: {type(exc).__name__}") from exc
        if len(response.content) < 100_000 or response.content[:2] != b"PK":
            raise GlobalProviderError("HKEX securities workbook payload is incomplete or not XLSX")
        return response.content

    def list_instruments(self, *, country: Optional[str] = None, exchange: Optional[str] = None):
        if not country or not exchange or not self.supported(country, exchange):
            raise GlobalProviderError(f"HKEX official universe is not configured for {country}/{exchange}")
        try:
            # HKEX currently publishes an incorrect worksheet dimension for this
            # workbook; normal mode intentionally ignores that read-only shortcut.
            workbook = openpyxl.load_workbook(io.BytesIO(self._download()), read_only=False, data_only=True)
        except Exception as exc:
            raise GlobalProviderError(f"HKEX securities workbook parse failed: {type(exc).__name__}") from exc
        worksheet = workbook["ListOfSecurities"] if "ListOfSecurities" in workbook.sheetnames else workbook.active
        rows = list(worksheet.iter_rows(values_only=True))
        result = parse_hkex_securities_rows(rows)
        self.last_metadata = {
            "officialCount": len(result),
            "resolvedCount": len(result),
            "resolutionCoveragePct": 100.0,
            "publicationDate": datetime.now(timezone.utc).date().isoformat(),
            "nativeMic": "XHKG",
            "identitySource": "HKEX official List of Securities workbook",
            "eligibleScope": "Main Board/GEM primary HKD equity counters",
            "sourceUrl": self.source_url,
        }
        return result
