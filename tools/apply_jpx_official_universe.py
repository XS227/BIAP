from pathlib import Path

PROVIDER = r'''"""Official Tokyo Stock Exchange ordinary-equity universe from JPX.

JPX's monthly "List of TSE-listed Issues" is the membership source. BIAP keeps
only domestic Prime/Standard/Growth listings with a standard four-character TSE
local code. This excludes ETFs/ETNs, REIT/fund products, PRO Market, foreign
listings and five-digit preferred/bond-type class shares.
"""
from __future__ import annotations

from datetime import datetime, timezone
import io
import re
from typing import Optional

import openpyxl
import requests

from .models import GlobalCompany, SourceEvidence
from .providers import GlobalProviderError, InstrumentUniverseProvider

_LIST_URL = (
    "https://www.jpx.co.jp/english/markets/statistics-equities/misc/"
    "tvdivq0000001vg2-att/data_e.xlsx"
)
_PROVIDER = "official-jpx-tse-listed-issues"
_DOMESTIC = {
    "primemarket(domestic)",
    "standardmarket(domestic)",
    "growthmarket(domestic)",
}
_CODE_RE = re.compile(r"^[0-9A-Z]{4}$")


def _text(value: object) -> str:
    return str(value or "").strip()


def _section_key(value: object) -> str:
    return re.sub(r"\s+", "", _text(value)).lower()


def _code(value: object) -> str:
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return _text(value).upper()


def is_jpx_domestic_common(code: object, section: object) -> bool:
    return bool(_CODE_RE.fullmatch(_code(code))) and _section_key(section) in _DOMESTIC


def parse_jpx_workbook(data: bytes) -> tuple[str, list[dict]]:
    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        ws = wb.active
        rows = ws.iter_rows(values_only=True)
        header = [str(value or "").strip() for value in next(rows)]
    except Exception as exc:
        raise GlobalProviderError(f"JPX listed-issues workbook is invalid: {type(exc).__name__}") from exc

    index = {name: i for i, name in enumerate(header)}
    required = {"Effective Date", "Local Code", "Name (English)", "Section/Products"}
    if not required.issubset(index):
        raise GlobalProviderError("JPX listed-issues workbook is missing required columns")

    effective = ""
    parsed: list[dict] = []
    seen: set[str] = set()
    for row in rows:
        code = _code(row[index["Local Code"]])
        section = _text(row[index["Section/Products"]])
        if not is_jpx_domestic_common(code, section) or code in seen:
            continue
        name = _text(row[index["Name (English)"]])
        if not name:
            continue
        date_value = _text(row[index["Effective Date"]])
        effective = max(effective, date_value)
        parsed.append({
            "code": code,
            "name": name,
            "section": section,
            "effectiveDate": date_value,
            "sector33Code": _text(row[index["33 Sector(Code)"]]) if "33 Sector(Code)" in index else None,
            "sector33Name": _text(row[index["33 Sector(name)"]]) if "33 Sector(name)" in index else None,
            "sector17Code": _text(row[index["17 Sector(Code)"]]) if "17 Sector(Code)" in index else None,
            "sector17Name": _text(row[index["17 Sector(name)"]]) if "17 Sector(name)" in index else None,
            "sizeCode": _text(row[index["Size Code (New Index Series)"]]) if "Size Code (New Index Series)" in index else None,
            "sizeName": _text(row[index["Size (New Index Series)"]]) if "Size (New Index Series)" in index else None,
        })
        seen.add(code)
    if not parsed:
        raise GlobalProviderError("JPX listed-issues workbook returned no domestic ordinary equities")
    return effective, parsed


class JPXOfficialUniverseProvider(InstrumentUniverseProvider):
    provider_id = _PROVIDER

    def __init__(self, *, timeout: float = 45.0) -> None:
        self.timeout = max(10.0, float(timeout))
        self.last_metadata: dict = {}

    def _download(self) -> bytes:
        try:
            response = requests.get(
                _LIST_URL,
                headers={
                    "User-Agent": "BIAP Global market-data integration (+https://setai.no)",
                    "Accept": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet,*/*",
                },
                timeout=self.timeout,
            )
            response.raise_for_status()
        except requests.RequestException as exc:
            raise GlobalProviderError(f"JPX listed-issues request failed: {type(exc).__name__}") from exc
        if not response.content.startswith(b"PK"):
            raise GlobalProviderError("JPX listed-issues response is not an XLSX workbook")
        return response.content

    def list_instruments(self, *, country: Optional[str] = None, exchange: Optional[str] = None):
        if str(country or "").upper() != "JP" or str(exchange or "").upper() != "TSE_JP":
            raise GlobalProviderError(f"JPX universe is not configured for {country}/{exchange}")
        effective, rows = parse_jpx_workbook(self._download())
        observed = datetime.now(timezone.utc).isoformat()
        result: list[GlobalCompany] = []
        for row in rows:
            result.append(GlobalCompany(
                country="JP",
                exchange="TSE_JP",
                currency="JPY",
                ticker=row["code"],
                name=row["name"],
                mic_code="XJPX",
                instrument_type="Common Stock",
                sector=row.get("sector17Name") or None,
                industry=row.get("sector33Name") or None,
                raw_provider_fields={
                    "official_universe": True,
                    "jpx_effective_date": row.get("effectiveDate"),
                    "jpx_section": row.get("section"),
                    "jpx_sector_33_code": row.get("sector33Code"),
                    "jpx_sector_17_code": row.get("sector17Code"),
                    "jpx_size_code": row.get("sizeCode"),
                    "jpx_size_name": row.get("sizeName"),
                },
                sources=[SourceEvidence(
                    provider=self.provider_id,
                    source_type="official_exchange_reference",
                    source_id=f"JPX-TSE-LIST-{effective or 'current'}",
                    source_url=_LIST_URL,
                    observed_at=observed,
                    period_end=effective or None,
                    quality=0.99,
                    notes="JPX List of TSE-listed Issues; domestic Prime/Standard/Growth ordinary-equity scope.",
                )],
            ))
        self.last_metadata = {
            "officialCount": len(result),
            "resolvedCount": len(result),
            "resolutionCoveragePct": 100.0,
            "effectiveDate": effective,
            "identitySource": "JPX List of TSE-listed Issues",
            "filter": "Prime/Standard/Growth domestic + standard 4-character TSE local code",
        }
        return result
'''

TEST = r'''import io

import openpyxl

from global_markets.jpx_official import is_jpx_domestic_common, parse_jpx_workbook


def test_jpx_domestic_filter_excludes_products_foreign_and_class_shares():
    assert is_jpx_domestic_common("1301", "Prime Market (Domestic)")
    assert is_jpx_domestic_common("130A", "Growth Market(Domestic)")
    assert is_jpx_domestic_common(1332, "Prime Market (Domestic)")
    assert not is_jpx_domestic_common("1305", "ETFs/ ETNs")
    assert not is_jpx_domestic_common("25935", "Prime Market (Domestic)")
    assert not is_jpx_domestic_common("50765", "Prime Market (Domestic)")
    assert not is_jpx_domestic_common("9999", "Standard Market(Foreign)")
    assert not is_jpx_domestic_common("131A", "PRO Market")


def test_parse_jpx_workbook_keeps_only_domestic_ordinary_equities():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append([
        "Effective Date", "Local Code", "Name (English)", "Section/Products",
        "33 Sector(Code)", "33 Sector(name)", "17 Sector(Code)", "17 Sector(name)",
        "Size Code (New Index Series)", "Size (New Index Series)",
    ])
    ws.append(["20260831", 1301, "KYOKUYO", "Prime Market (Domestic)", "50", "Fishery", "1", "FOODS", "6", "Small"])
    ws.append(["20260831", "130A", "Veritas", "Growth Market(Domestic)", "3250", "Pharmaceutical", "5", "PHARMA", "-", "-"])
    ws.append(["20260831", 1305, "ETF", "ETFs/ ETNs", "-", "-", "-", "-", "-", "-"])
    ws.append(["20260831", 25935, "Preferred", "Prime Market (Domestic)", "3050", "Foods", "1", "FOODS", "-", "-"])
    out = io.BytesIO()
    wb.save(out)
    effective, rows = parse_jpx_workbook(out.getvalue())
    assert effective == "20260831"
    assert [row["code"] for row in rows] == ["1301", "130A"]
'''

Path("analysis/global_markets/jpx_official.py").write_text(PROVIDER, encoding="utf-8")
Path("analysis/tests/test_global_jpx_official.py").write_text(TEST, encoding="utf-8")

p=Path("analysis/global_markets/runtime.py")
text=p.read_text(encoding="utf-8")
anchor='from .iran_adapter import IranLegacyProvider\n'
insert=anchor+'from .jpx_official import JPXOfficialUniverseProvider\n'
if insert not in text:
    if anchor not in text: raise SystemExit("runtime import anchor missing")
    text=text.replace(anchor,insert,1)
anchor='''    asx_universe = PersistentUniverseProvider(ASXUniverseProvider())
    registry.register_universe("AU", "ASX", asx_universe)

'''
insert=anchor+'''    jpx_universe = PersistentUniverseProvider(
        JPXOfficialUniverseProvider(),
        fresh_hours=24,
    )
    registry.register_universe("JP", "TSE_JP", jpx_universe)

'''
if insert not in text:
    if anchor not in text: raise SystemExit("runtime registration anchor missing")
    text=text.replace(anchor,insert,1)
p.write_text(text,encoding="utf-8")

p=Path("analysis/global_markets/cached_universe.py")
text=p.read_text(encoding="utf-8")
old='''# Version 15 replaces the UK reference/demo catalog with the official London
# Stock Exchange Main Market universe. Older LSE snapshots must not survive the
# authoritative source change.
CACHE_SCHEMA_VERSION = 15
'''
new='''# Version 16 replaces the Japan reference/demo catalog with the official JPX
# TSE domestic ordinary-equity universe. Older Japan snapshots must not survive
# the authoritative source change.
CACHE_SCHEMA_VERSION = 16
'''
if old not in text: raise SystemExit("cache schema anchor missing")
p.write_text(text.replace(old,new,1),encoding="utf-8")

p=Path("analysis/tests/test_global_universe_cache_v8.py")
text=p.read_text(encoding="utf-8")
old='''def test_universe_cache_schema_is_v15():
    assert CACHE_SCHEMA_VERSION == 15
'''
new='''def test_universe_cache_schema_is_v16():
    assert CACHE_SCHEMA_VERSION == 16
'''
if old not in text: raise SystemExit("cache test anchor missing")
p.write_text(text.replace(old,new,1),encoding="utf-8")
