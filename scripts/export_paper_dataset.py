#!/usr/bin/env python3
"""Export ~/biap_research/paper_dataset.xlsx from the FROZEN snapshot (one row per issuer, 10 countries).

Sheets: Stocks (per stock), Gate_Cascade (per country: stocks removed by each conservative gate, in order),
Tier_Counts (per country, both modes), Notes. Reads only the snapshot; the manifest is verified first.

    python3 scripts/export_paper_dataset.py [--snapshot ~/biap_research/snapshot_2026-10-19] [--out ~/biap_research/paper_dataset.xlsx]
"""
from __future__ import annotations

import argparse
import sys
from collections import Counter, defaultdict
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import research_common as C  # noqa: E402

A = C.A
GROUPS = [name for name, _ in C.CASCADE]


def liquidity_ranks(rows, counted) -> dict:
    """Rank 1 = most traded inside the market (GB: market cap, because GB liquidityValue is market cap)."""
    by: dict = defaultdict(list)
    for r in rows:
        if r.analysed and r.key in counted:
            by[C.market_of(r.exchange)].append(r)
    out = {}
    for rs in by.values():
        ordered = sorted((r for r in rs if r.liquidity > 0), key=lambda r: -r.liquidity)
        for i, r in enumerate(ordered, 1):
            out[r.key] = i
    return out


def stock_rows(rows, decisions, counted):
    ranks = liquidity_ranks(rows, counted)
    for r in sorted((r for r in rows if r.analysed and r.key in counted), key=lambda r: (r.country, C.market_of(r.exchange), r.ticker)):
        d = decisions[r.key]
        failed, _ = A.failed_gates(r.gate, C.GATE_PARAMS)
        line = [r.country, C.market_of(r.exchange), r.exchange, r.ticker, r.name, r.isin or "", r.sector or ""]
        for a in C.AGENTS:
            s = r.candidate.signals.get(a)
            line += [None, None] if s is None else [round(s[0], 3), round(s[1], 3)]
        line += [d.agents_with_data, d.score, d.confidence, r.gate.score, r.gate.confidence, r.stored_call, C.gate_group(failed), "; ".join(failed),
                 d.call, d.tier_before_flags or "", ",".join(d.flags), d.peer_scope, ranks.get(r.key),
                 round(r.liquidity) if r.liquidity else None, "sector_reference" if r.sector_filled else ""]
        yield line, r, failed


def main(argv=None) -> int:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("--snapshot", default=str(C.DEFAULT_SNAPSHOT))
    ap.add_argument("--tmp-dir", default=str(C.DEFAULT_TMP))
    ap.add_argument("--out", default="~/biap_research/paper_dataset.xlsx")
    args = ap.parse_args(argv)
    rows, decisions, counted = C.load(Path(args.snapshot), args.tmp_dir)

    wb = Workbook()
    wb.remove(wb.active)
    bold, head = Font(bold=True), PatternFill("solid", fgColor="DDE6F0")

    def sheet(title, header, data, width=14):
        ws = wb.create_sheet(title)
        ws.append(header)
        for c in ws[1]:
            c.font, c.fill = bold, head
        for line in data:
            ws.append(line)
        ws.freeze_panes = "A2"
        for i in range(1, len(header) + 1):
            ws.column_dimensions[get_column_letter(i)].width = width
        return ws

    header = ["country", "market", "listing", "ticker", "name", "ISIN", "sector"]
    for a in C.AGENTS:
        header += [f"{a}_vote", f"{a}_conf"]
    header += ["agents_with_data", "inv_score", "inv_confidence", "conservative_score", "conservative_confidence", "conservative_call", "conservative_primary_gate",
               "conservative_gates_failed", "investment_tier", "tier_before_flags", "flags", "peer_scope",
               "liquidity_rank", "liquidity_value", "note"]
    cascade: dict = defaultdict(Counter)
    stored_buy: Counter = Counter()
    data = []
    for line, r, failed in stock_rows(rows, decisions, counted):
        data.append(line)
        cascade[r.country]["analysed"] += 1
        cascade[r.country][line[header.index("conservative_primary_gate")]] += 1
        stored_buy[r.country] += r.stored_call == "BUY_CANDIDATE"
    ws = sheet("Stocks", header, data)
    ws.column_dimensions["E"].width = 30
    ws.column_dimensions["Z"].width = 40

    crow = []
    for cc in sorted(cascade):
        c = cascade[cc]
        crow.append([cc, c["analysed"]] + [c[g] for g in GROUPS] + [c["PASS"], stored_buy[cc]])
    sheet("Gate_Cascade", ["country", "analysed"] + [f"removed_by_{g}" for g in GROUPS] +
          ["pass_all_gates (BUY under PR52 rules)", "stored_conservative_BUY"], crow, 18)
    ws = wb["Gate_Cascade"]
    ws.append([])
    ws.append(["remaining after each gate"])
    ws.append(["country", "analysed"] + [f"after_{g}" for g in GROUPS])
    for cc in sorted(cascade):
        left, line = cascade[cc]["analysed"], [cc, cascade[cc]["analysed"]]
        for g in GROUPS:
            left -= cascade[cc][g]
            line.append(left)
        ws.append(line)

    counts = C.tier_counts(rows, decisions, counted, "country")
    sheet("Tier_Counts", ["country", "analysed"] + [f"conservative_{t}" for t in C.CONS_TIERS] + [f"investment_{t}" for t in C.INV_TIERS],
          [[cc, v["analysed"]] + [v["conservative"][t] for t in C.CONS_TIERS] + [v["investment"][t] for t in C.INV_TIERS]
           for cc, v in counts.items()], 16)

    ws = wb.create_sheet("Notes")
    for line in [
        f"Source: frozen snapshot {Path(args.snapshot).expanduser()} (sha256 manifest verified). Scope: INVESTMENT_MODE_MARKETS, 10 countries.",
        "One row per issuer: Xetra and Frankfurt are merged by ISIN (DE); US = NASDAQ + NYSE ranked as two markets, shown as country US.",
        "conservative_call = the call stored by the conservative engine; investment_tier = investment mode (BUY needs score>0, >=4 agents with data, top 10% of its market).",
        "Agent vote/confidence are the raw stored agent outputs (0 / 0 = the agent had no data; blank = no signal stored). agents_with_data counts agents with confidence >= 0.30 (investment-mode rule).",
        "inv_score / inv_confidence come from investment mode (unknown agents left out); conservative_score / conservative_confidence are the stored conservative-engine values the gates test.",
        "conservative_primary_gate = the FIRST failed group in the order: " + ", ".join(GROUPS) + "; PASS = no gate failed.",
        "The cascade evaluates the production gates PLUS the PR #52 BUY rules (distress not LOW_RISK, peers >= 8, independent positive signal, comparison evidence). "
        "Stored calls pre-date those rules, so a stored BUY can fail one of them here (compare pass_all_gates with stored_conservative_BUY).",
        "liquidity_rank: 1 = most traded in the market (quote.liquidityValue; GB = market cap). Blank when the ledger has no liquidity value.",
        "note = sector_reference when the sector was filled from the ISIN-keyed sector reference (replay only).",
    ]:
        ws.append([line])
    ws.column_dimensions["A"].width = 160
    out = Path(args.out).expanduser()
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print(f"wrote {out} ({len(data)} stocks)")
    for cc, v in counts.items():
        print(cc, v["analysed"], "C", list(v["conservative"].values()), "I", list(v["investment"].values()))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
